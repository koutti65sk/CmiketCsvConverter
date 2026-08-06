import pandas as pd
from openpyxl import load_workbook


def get_hyperlinks(ws, header_name):
    """
    指定した列のハイパーリンク一覧を取得する
    """

    header = {
        cell.value: i + 1
        for i, cell in enumerate(ws[1])
    }

    if header_name not in header:
        return []

    col = header[header_name]

    links = []

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row, col)

        if cell.hyperlink:
            links.append(cell.hyperlink.target)
        else:
            links.append("")

    return links


def read_excel(file_path, max_rows=None):
    """
    Excelを読み込み、
    シートごとのデータを返す
    """

    excel = pd.ExcelFile(file_path)

    wb = load_workbook(
        file_path,
        data_only=True
    )

    sheets = []

    for sheet_name in excel.sheet_names:

        df = pd.read_excel(
            file_path,
            sheet_name=sheet_name,
            nrows=max_rows
        )

        # 列名の余計な空白削除
        df.columns = (
            df.columns
            .astype(str)
            .str.strip()
        )

        sheets.append(
            {
                "name": sheet_name,
                "df": df,
                "ws": wb[sheet_name]
            }
        )

    return sheets